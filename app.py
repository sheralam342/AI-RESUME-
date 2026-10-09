"""AI Resume ATS Checker - Streamlit + Groq.

Upload a resume (PDF or DOCX), optionally paste a job description, and get an
ATS score with concrete suggestions for improvement.
"""

import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "openai/gpt-oss-120b"
MODEL_OPTIONS = [
    "openai/gpt-oss-120b",
]
MAX_RESUME_CHARS = 12000  # keep prompts within model limits
MAX_JD_CHARS = 6000
MIN_RESUME_CHARS = 200  # below this the PDF is probably scanned/empty
MAX_FILE_MB = 5

STANDARD_SECTIONS = {
    "Experience": r"\b(experience|employment|work history)\b",
    "Education": r"\beducation\b",
    "Skills": r"\bskills\b",
    "Summary/Objective": r"\b(summary|objective|profile)\b",
    "Projects": r"\bprojects?\b",
}

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume strictly and honestly. \
Do not inflate scores. Respond with ONLY a valid JSON object, no markdown, \
no commentary."""

JSON_SCHEMA_HINT = """Return a JSON object with exactly these keys:
{
  "ats_score": <integer 0-100, overall ATS compatibility and quality>,
  "score_breakdown": {
    "keywords": <integer 0-100>,
    "formatting": <integer 0-100>,
    "content_quality": <integer 0-100>,
    "impact_and_metrics": <integer 0-100>
  },
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<short point>", ...],
  "missing_keywords": ["<keyword or skill>", ...],
  "improvements": [
    {"area": "<section or topic>", "issue": "<what is wrong>", "fix": "<specific actionable fix>"}
  ],
  "formatting_issues": ["<issue>", ...]
}
Give 4-8 improvements ordered by importance. If no job description is given, \
judge keywords against common expectations for the candidate's apparent target role."""


# ----------------------------------------------------------------------------
# File parsing
# ----------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password-protected.")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Resumes often put content in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ----------------------------------------------------------------------------
# Rule-based checks (fast, deterministic, no AI needed)
# ----------------------------------------------------------------------------
def run_basic_checks(text: str) -> dict:
    lower = text.lower()
    words = re.findall(r"\b\w+\b", text)
    return {
        "email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
        "phone": bool(re.search(r"(\+?\d[\d\s().-]{7,}\d)", text)),
        "linkedin": "linkedin.com" in lower,
        "word_count": len(words),
        "sections": {
            label: bool(re.search(pattern, lower))
            for label, pattern in STANDARD_SECTIONS.items()
        },
        "has_numbers": bool(re.search(r"\d+\s*(%|\+|k\b|x\b)|\$\s*\d", lower)),
    }


# ----------------------------------------------------------------------------
# AI analysis
# ----------------------------------------------------------------------------
def build_user_prompt(resume_text: str, job_description: str) -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    prompt = f"{JSON_SCHEMA_HINT}\n\n=== RESUME ===\n{resume_text}\n"
    if job_description.strip():
        prompt += f"\n=== JOB DESCRIPTION ===\n{job_description.strip()[:MAX_JD_CHARS]}\n"
    else:
        prompt += "\n(No job description provided.)\n"
    return prompt


def parse_json_response(content: str) -> dict:
    """Parse model output as JSON, tolerating stray text or code fences."""
    content = content.strip()
    content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content, flags=re.IGNORECASE)
    try:
        return json.loads(content)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", content, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def _clamp_score(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _as_str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if v is not None and str(v).strip()]


def normalize_result(raw: dict) -> dict:
    """Make sure the model output has every field the UI expects."""
    breakdown = raw.get("score_breakdown") if isinstance(raw.get("score_breakdown"), dict) else {}
    improvements = []
    for item in raw.get("improvements", []) if isinstance(raw.get("improvements"), list) else []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "area": str(item.get("area", "General")),
                    "issue": str(item.get("issue", "")),
                    "fix": str(item.get("fix", "")),
                }
            )
        elif isinstance(item, str):
            improvements.append({"area": "General", "issue": "", "fix": item})
    return {
        "ats_score": _clamp_score(raw.get("ats_score")),
        "score_breakdown": {
            "Keywords": _clamp_score(breakdown.get("keywords")),
            "Formatting": _clamp_score(breakdown.get("formatting")),
            "Content quality": _clamp_score(breakdown.get("content_quality")),
            "Impact & metrics": _clamp_score(breakdown.get("impact_and_metrics")),
        },
        "summary": str(raw.get("summary", "")),
        "strengths": _as_str_list(raw.get("strengths")),
        "missing_keywords": _as_str_list(raw.get("missing_keywords")),
        "improvements": improvements,
        "formatting_issues": _as_str_list(raw.get("formatting_issues")),
    }


def analyze_resume(client, model: str, resume_text: str, job_description: str) -> dict:
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": build_user_prompt(resume_text, job_description)},
        ],
        temperature=0.2,
        max_tokens=2048,
        response_format={"type": "json_object"},
    )
    content = response.choices[0].message.content or ""
    return normalize_result(parse_json_response(content))


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def get_api_key() -> str:
    try:
        key = st.secrets.get("GROQ_API_KEY", "")
    except Exception:  # no secrets file present
        key = ""
    return key or os.environ.get("GROQ_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_basic_checks(checks: dict) -> None:
    st.subheader("Quick checks")
    cols = st.columns(3)
    cols[0].metric("Email", "Found" if checks["email"] else "Missing")
    cols[1].metric("Phone", "Found" if checks["phone"] else "Missing")
    cols[2].metric("LinkedIn", "Found" if checks["linkedin"] else "Missing")

    wc = checks["word_count"]
    if wc < 250:
        st.warning(f"Resume is short ({wc} words). Aim for roughly 400-800.")
    elif wc > 1000:
        st.warning(f"Resume is long ({wc} words). Aim for 1-2 pages.")
    else:
        st.success(f"Length looks reasonable ({wc} words).")

    missing = [name for name, ok in checks["sections"].items() if not ok]
    if missing:
        st.info("Sections not detected: " + ", ".join(missing))
    if not checks["has_numbers"]:
        st.info("No quantified achievements detected (%, $, numbers). Add metrics.")


def render_results(result: dict) -> None:
    score = result["ats_score"]
    st.header(f"ATS Score: {score}/100 - {score_label(score)}")
    st.progress(score / 100)
    if result["summary"]:
        st.write(result["summary"])

    cols = st.columns(len(result["score_breakdown"]))
    for col, (label, value) in zip(cols, result["score_breakdown"].items()):
        col.metric(label, f"{value}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for item in result["strengths"] or ["None identified."]:
            st.markdown(f"- {item}")
    with right:
        st.subheader("Missing keywords")
        if result["missing_keywords"]:
            st.markdown(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.markdown("None identified.")

    st.subheader("Suggested improvements")
    for i, imp in enumerate(result["improvements"], start=1):
        with st.expander(f"{i}. {imp['area']}", expanded=i <= 3):
            if imp["issue"]:
                st.markdown(f"**Issue:** {imp['issue']}")
            st.markdown(f"**Fix:** {imp['fix']}")

    if result["formatting_issues"]:
        st.subheader("Formatting issues")
        for item in result["formatting_issues"]:
            st.markdown(f"- {item}")


# ----------------------------------------------------------------------------
# App
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and specific ways to improve it.")

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Groq API key", type="password",
                                    help="Get a free key at console.groq.com")
        model = st.selectbox("Model", MODEL_OPTIONS, index=0)
        st.markdown("---")
        st.caption("Your resume is sent to Groq for analysis and is not stored by this app.")

    col_a, col_b = st.columns(2)
    with col_a:
        uploaded = st.file_uploader("Resume (PDF or DOCX)", type=["pdf", "docx"])
    with col_b:
        job_description = st.text_area(
            "Job description (optional, improves keyword matching)", height=180
        )

    if not st.button("Analyze resume", type="primary"):
        return

    if not api_key:
        st.error("Please provide a Groq API key in the sidebar.")
        return
    if uploaded is None:
        st.error("Please upload a resume first.")
        return
    if uploaded.size > MAX_FILE_MB * 1024 * 1024:
        st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
        return

    try:
        text = extract_resume_text(uploaded.name, uploaded.getvalue())
    except Exception as exc:
        st.error(f"Could not read the file: {exc}")
        return

    if len(text) < MIN_RESUME_CHARS:
        st.error(
            "Very little text could be extracted. If your resume is a scanned image, "
            "that is also a problem for real ATS systems - export a text-based PDF or DOCX."
        )
        return

    with st.spinner("Analyzing your resume..."):
        try:
            result = analyze_resume(Groq(api_key=api_key), model, text, job_description)
        except json.JSONDecodeError:
            st.error("The model returned an unreadable response. Please try again.")
            return
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            return

    render_basic_checks(run_basic_checks(text))
    st.divider()
    render_results(result)
    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


if __name__ == "__main__":
    main()
